# -*- coding: utf-8 -*-
from pathlib import Path
from zipfile import ZipFile
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
import json
root = Path('artifacts/native')
for name in ['sample.docx', 'sample.xlsx']:
    with ZipFile(root/name) as archive:
        assert archive.testzip() is None
word = Document(root/'sample.docx')
assert any('设备采购' in p.text for p in word.paragraphs)
assert word.tables[0].cell(1, 0).text == '27 英寸显示器'
book = load_workbook(root/'sample.xlsx')
assert book.active['A2'].value == '27 英寸显示器'
assert book.active['D5'].value == '5,113'
assert book.active['B2'].data_type == 's'
pdf = PdfReader(root/'sample.pdf')
assert len(pdf.pages) == 1
assert len(pdf.pages[0].images) >= 1
ocr = json.loads((root/'ocr-result.json').read_text())
assert ocr['engine'] == 'Apple Vision'
assert '设备' in ocr['text'] and len(ocr['blocks']) > 10
print(json.dumps({'docx': 'valid editable paragraphs and table', 'xlsx': 'valid string cells and table', 'pdf': 'valid one-page image PDF', 'ocr': {'engine': ocr['engine'], 'blocks': len(ocr['blocks'])}}, ensure_ascii=False, indent=2))
